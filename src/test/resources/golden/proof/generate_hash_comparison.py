#!/usr/bin/env python3
"""Before-vs-after hash comparison of the golden (legacy-backed) API responses
against the live (modern-repository-backed) application.

Usage: python3 generate_hash_comparison.py [BASE_URL]   (default http://localhost:8080)

"Before" = the C0 golden files in src/test/resources/golden/.
"After"  = live GET responses from the running app for the same 17 endpoints.

Normalization (applied identically to both sides before hashing):
    json.dumps(json.loads(text), sort_keys=True, indent=4) + "\\n"
i.e. exactly the output of `python3 -m json.tool --sort-keys`.
Raw SHA-256 of the untouched golden file bytes and of the untouched live body
are also recorded so whitespace-only differences are visible but not fatal.

Requires: openpyxl (pip install openpyxl). Writes golden_hash_comparison.xlsx
next to this script.
"""
import hashlib
import json
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8080"
HERE = Path(__file__).resolve().parent
GOLDEN_DIR = HERE.parent
OUT = HERE / "golden_hash_comparison.xlsx"
NORMALIZATION = ("json.dumps(json.loads(text), sort_keys=True, indent=4) + '\\n' "
                 "(== `python3 -m json.tool --sort-keys`)")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize(text: str) -> str:
    return json.dumps(json.loads(text), sort_keys=True, indent=4) + "\n"


def fetch(path: str) -> bytes:
    with urllib.request.urlopen(BASE_URL + path, timeout=30) as resp:
        return resp.read()


def endpoints() -> list[tuple[str, str]]:
    loans = json.loads((GOLDEN_DIR / "loans.json").read_text())
    borrowers = json.loads((GOLDEN_DIR / "borrowers.json").read_text())
    eps = [("/api/loans", "loans.json"), ("/api/borrowers", "borrowers.json")]
    for loan in loans:
        lid = loan["loanAccountNumber"]
        eps.append((f"/api/loans/{lid}", f"loan_{lid}.json"))
        eps.append((f"/api/loans/{lid}/payments", f"payments_{lid}.json"))
    for b in borrowers:
        eps.append((f"/api/borrowers/{b['id']}", f"borrower_{b['id']}.json"))
    return eps


def main() -> int:
    rows = []
    for path, golden_name in endpoints():
        golden_raw = (GOLDEN_DIR / golden_name).read_bytes()
        live_raw = fetch(path)
        before_norm = normalize(golden_raw.decode())
        after_norm = normalize(live_raw.decode())
        h_before, h_after = sha256(before_norm.encode()), sha256(after_norm.encode())
        rows.append({
            "endpoint": path,
            "golden_file": golden_name,
            "normalized_hash_before": h_before,
            "normalized_hash_after": h_after,
            "raw_hash_before": sha256(golden_raw),
            "raw_hash_after": sha256(live_raw),
            "raw_bytes_before": len(golden_raw),
            "raw_bytes_after": len(live_raw),
            "result": "MATCH" if h_before == h_after else "MISMATCH",
        })

    matches = sum(r["result"] == "MATCH" for r in rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "hash_comparison"
    headers = ["#", "Endpoint", "Golden file (before)", "Normalized SHA-256 (before)",
               "Normalized SHA-256 (after)", "Raw SHA-256 before (golden file)",
               "Raw SHA-256 after (live body)", "Raw bytes before", "Raw bytes after",
               "Result"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for i, r in enumerate(rows, 1):
        ws.append([i, r["endpoint"], r["golden_file"], r["normalized_hash_before"],
                   r["normalized_hash_after"], r["raw_hash_before"], r["raw_hash_after"],
                   r["raw_bytes_before"], r["raw_bytes_after"], r["result"]])
        ws.cell(row=ws.max_row, column=10).fill = green if r["result"] == "MATCH" else red

    ws.append([])
    summary = f"SUMMARY: {matches}/{len(rows)} endpoints MATCH"
    ws.append(["", summary, "", "", "", "", "", "", "",
               "MATCH" if matches == len(rows) else "MISMATCH"])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=10).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=10).fill = green if matches == len(rows) else red

    meta = wb.create_sheet("method")
    meta.append(["Generated (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    meta.append(["Base URL (after)", BASE_URL])
    meta.append(["Golden directory (before)", str(GOLDEN_DIR.relative_to(HERE.parents[3]))])
    meta.append(["Normalization", NORMALIZATION])
    meta.append(["Hash", "SHA-256 (hashlib) over UTF-8 bytes"])
    meta.append(["Comparison rule", "MATCH iff normalized SHA-256 before == normalized SHA-256 after"])
    meta.append(["Raw hashes", "Informational: golden files are pretty-printed with indent 2 "
                               "while the live body is compact JSON, so raw hashes differ by "
                               "whitespace only when the normalized hashes match."])
    for sheet in (ws, meta):
        for col in range(1, sheet.max_column + 1):
            width = max(len(str(c.value or "")) for c in sheet[get_column_letter(col)])
            sheet.column_dimensions[get_column_letter(col)].width = min(width + 2, 70)

    wb.save(OUT)
    for r in rows:
        print(f"{r['result']:8} {r['endpoint']}")
    print(summary)
    print(f"wrote {OUT}")
    return 0 if matches == len(rows) else 1


if __name__ == "__main__":
    sys.exit(main())
